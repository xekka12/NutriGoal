from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

GREEN2 = RGBColor(0x10, 0xB9, 0x81)
DARK = RGBColor(0x0B, 0x12, 0x20)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xE2, 0xE8, 0xF0)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
CARD = RGBColor(0x0F, 0x17, 0x2A)


def set_run(run, size=18, bold=False, color=WHITE):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_bg(slide, color=DARK):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_accent_bar(slide):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN2
    bar.line.fill.background()


def add_title(slide, text, top=0.35, size=36):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(12), Inches(0.8))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=True, color=WHITE)


def add_subtitle(slide, text, top=1.05, size=18, color=MUTED):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(12), Inches(0.5))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, size=size, color=color)


def add_bullets(slide, items, left=0.7, top=1.7, width=12, size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(12)
        run = p.add_run()
        run.text = f"•  {item}"
        set_run(run, size=size, color=LIGHT)


def add_card(slide, left, top, width, height, title, body):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = GREEN2
    shape.line.width = Pt(1.5)

    tbox = slide.shapes.add_textbox(
        Inches(left + 0.25), Inches(top + 0.25), Inches(width - 0.5), Inches(0.45)
    )
    p = tbox.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run(run, size=18, bold=True, color=GREEN2)

    bbox = slide.shapes.add_textbox(
        Inches(left + 0.25), Inches(top + 0.75), Inches(width - 0.5), Inches(height - 1)
    )
    tf = bbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = body
    set_run(run, size=14, color=LIGHT)


# Slide 1
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_accent_bar(s)
add_title(s, "NutriGoal", top=2.2, size=48)
add_subtitle(s, "Week 2 Progress Report", top=3.1, size=28, color=GREEN2)
add_subtitle(
    s,
    "Smarter Filtering Experience: Health Status, Image Cards & Allergies",
    top=3.8,
    size=18,
)
add_subtitle(s, "Healthy Meal Recommender App", top=4.5, size=16)

# Slide 2
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_accent_bar(s)
add_title(s, "Week 2 Agenda")
add_subtitle(s, "What we improved this week")
add_bullets(
    s,
    [
        "What is the focus of Week 2?",
        "New Health Status & Illness filter",
        "Image Card Picker (replacing dropdowns)",
        "More Activity & Body Goal options",
        "Allergy Smart Cards",
        "User experience improvements",
        "Summary & next steps",
    ],
)

# Slide 3
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_accent_bar(s)
add_title(s, "Week 2 Focus")
add_subtitle(s, "Making Personalize Your Plan easier to understand and use")
add_bullets(
    s,
    [
        "Problem before: filters were plain dropdowns — hard to visualize choices",
        "Goal: help users pick options faster with images, names, and short descriptions",
        "Also added health-aware filtering: does the user have an illness or not?",
        "If yes, they can specify the condition (e.g., Heart Disease, Diabetes)",
        "Allergies were upgraded to the same smart card style",
    ],
)

# Slide 4
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_accent_bar(s)
add_title(s, "New Filter: Health Status")
add_subtitle(s, "Does the user currently have an illness or medical condition?")
add_card(
    s,
    0.7,
    1.8,
    5.7,
    3.8,
    "No Illness",
    "User is currently healthy.\nNo medical condition to consider.\nMeal recommendations stay general and balanced.",
)
add_card(
    s,
    6.9,
    1.8,
    5.7,
    3.8,
    "Has Illness / Medical Condition",
    "User currently has a condition.\nA second filter appears to specify what it is.\nRecommendations can lean toward lighter / more suitable meals.",
)

# Slide 5
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_accent_bar(s)
add_title(s, "Specify the Illness / Condition")
add_subtitle(s, "Shown only when “Has Illness” is selected")
add_bullets(
    s,
    [
        "Heart Disease",
        "High Blood Pressure (Hypertension)",
        "Diabetes",
        "Kidney Disease / Liver Disease",
        "Gout, Asthma, PCOS",
        "Plus: Specify other… for custom conditions",
        "This detail is carried into category matching and meal generation",
    ],
)

# Slide 6
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_accent_bar(s)
add_title(s, "From Dropdown → Image Card Picker")
add_subtitle(s, "A more visual way to choose filters")
add_card(
    s,
    0.7,
    1.8,
    3.8,
    4.2,
    "Before",
    "• Plain dropdown menus\n• Text-only options\n• Harder for users to understand\n• Less engaging selection",
)
add_card(
    s,
    4.8,
    1.8,
    3.8,
    4.2,
    "After",
    "• Tap a field → modal opens\n• Image + name + short description\n• Easy to compare choices\n• Specify other still available",
)
add_card(
    s,
    8.9,
    1.8,
    3.8,
    4.2,
    "Applies To",
    "• What do you usually do?\n• Body Goal\n• Health Status\n• Illness / Condition\n• Allergies",
)

# Slide 7
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_accent_bar(s)
add_title(s, "Activity Options Expanded")
add_subtitle(s, "“What do you usually do?” — clearer choices with related images")
add_bullets(
    s,
    [
        "Sedentary — mostly sitting / desk lifestyle",
        "Lightly Active — walking & chores",
        "Moderately Active — exercise 3–4x per week",
        "Gym / Workout — training 5–6x per week",
        "Athlete — daily intense training (image fixed)",
        "Heavy Work — construction / warehouse jobs",
        "Senior / Low Mobility — gentle movement (senior-related image)",
    ],
)

# Slide 8
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_accent_bar(s)
add_title(s, "Body Goal Options Expanded")
add_subtitle(s, "More goals so users can match their real intention")
add_bullets(
    s,
    [
        "Maintain Weight",
        "Lose Weight (Cutting)",
        "Gain Weight (image fixed)",
        "Build Muscle (Bulking)",
        "Body Recomp — lose fat + build muscle",
        "More Energy / Stamina",
        "Tone Body",
        "Stay Healthy",
    ],
)

# Slide 9
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_accent_bar(s)
add_title(s, "Allergy Smart Cards")
add_subtitle(s, "Multi-select cards instead of a plain checklist only")
add_bullets(
    s,
    [
        "Each allergy has an image, name, and short description",
        "Examples: Dairy, Peanuts, Fish, Shellfish, Soy, Pork, Beef, Chicken",
        "Users can select multiple allergies at once",
        "Custom allergy input is still available (e.g., Mango, MSG)",
        "Selected allergies are avoided in meal recommendations",
    ],
)

# Slide 10
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_accent_bar(s)
add_title(s, "Why This Helps Users")
add_subtitle(s, "UX improvements delivered this week")
add_card(
    s,
    0.7,
    1.8,
    3.8,
    4.0,
    "Clearer Choices",
    "Images + short descriptions help users understand what each option means.",
)
add_card(
    s,
    4.8,
    1.8,
    3.8,
    4.0,
    "Health-Aware",
    "Meals can consider illness status and specific conditions like heart disease or diabetes.",
)
add_card(
    s,
    8.9,
    1.8,
    3.8,
    4.0,
    "Safer Meals",
    "Allergy cards make it easier to mark foods to avoid before generating a plan.",
)

# Slide 11
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_accent_bar(s)
add_title(s, "User Flow (Demo Path)")
add_subtitle(s, "How a user personalizes their plan now")
add_bullets(
    s,
    [
        "1. Open Dashboard → Personalize Your Plan",
        "2. Tap Activity → pick an image card (or Specify other)",
        "3. Tap Body Goal → choose goal card",
        "4. Tap Health Status → No Illness / Has Illness",
        "5. If Has Illness → choose condition (e.g., Heart Disease)",
        "6. Tap Allergies → multi-select smart cards",
        "7. Optional budget → Find Matching Meal Categories",
    ],
)

# Slide 12
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_accent_bar(s)
add_title(s, "Week 2 Summary")
add_subtitle(s, "What we accomplished")
add_bullets(
    s,
    [
        "Added Health Status + specific illness/condition filtering",
        "Replaced dropdown filters with image card pickers",
        "Expanded Activity and Body Goal options",
        "Upgraded Allergies to smart image cards (multi-select)",
        "Improved descriptions and related images for clearer UX",
        "Result: faster, clearer, more personalized meal planning setup",
    ],
)

# Slide 13
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_accent_bar(s)
add_title(s, "Next Steps")
add_subtitle(s, "Possible focus for the following weeks")
add_bullets(
    s,
    [
        "Further polish of recommendation explanations",
        "More testing with different health profiles",
        "Continue refining meal generation quality",
        "Prepare Week 3 demo / documentation updates",
    ],
)

# Slide 14
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_accent_bar(s)
add_title(s, "Thank You", top=2.6, size=48)
add_subtitle(s, "NutriGoal — Week 2 Progress", top=3.5, size=22, color=GREEN2)
add_subtitle(s, "Questions?", top=4.3, size=18)

out = Path(r"C:\NutriGoal\NutriGoal_Week2_Progress_Report.pptx")
prs.save(out)
print(f"Saved: {out}")
print(f"Slides: {len(prs.slides)}")
